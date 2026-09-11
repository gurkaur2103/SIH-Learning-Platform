from io import BytesIO
import pymupdf
from docx import Document
from pptx import Presentation
from app.utils.validators import clean_text

def parse_document(data: bytes, filename: str) -> dict:
    lower = filename.lower()
    chunks = []
    if lower.endswith(".pdf"):
        pdf = pymupdf.open(stream=data, filetype="pdf")
        for index, page in enumerate(pdf):
            text = clean_text(page.get_text("text"))
            if text: chunks.append({"id":f"page-{index+1}","page":index+1,"text":text})
    elif lower.endswith(".docx"):
        doc = Document(BytesIO(data))
        for index, paragraph in enumerate(doc.paragraphs):
            text = clean_text(paragraph.text)
            if text: chunks.append({"id":f"paragraph-{index+1}","text":text})
    elif lower.endswith(".pptx"):
        presentation = Presentation(BytesIO(data))
        for index, slide in enumerate(presentation.slides):
            text = clean_text("\n".join(shape.text for shape in slide.shapes
                              if hasattr(shape, "text") and shape.text))
            if text: chunks.append({"id":f"slide-{index+1}","slide":index+1,"text":text})
    elif lower.endswith(".txt"):
        text = clean_text(data.decode("utf-8", errors="replace"))
        chunks = [{"id":"text-1","text":text}]
    else:
        raise ValueError("Supported files: PDF, DOCX, PPTX and TXT")
    return {"filename":filename,"text":"\n\n".join(c["text"] for c in chunks),
            "chunks":chunks,"characters":sum(len(c["text"]) for c in chunks)}
